# AlchemyZ module search with text + visual extraction for PDF/PPTX.
# Place this file beside Zach.py and put your lesson files in ./Modules/
import re
from pathlib import Path

SCRIPT_FOLDER = Path(__file__).resolve().parent
MODULE_FOLDER = SCRIPT_FOLDER / "Modules"

def _resolve_module_folder():
    """Find the Modules folder beside this script or in the current project directory."""
    candidates = [
        MODULE_FOLDER,
        Path.cwd() / "Modules",
        SCRIPT_FOLDER.parent / "Modules",
    ]
    for candidate in candidates:
        if candidate.is_dir():
            return candidate
    return MODULE_FOLDER
STOP_WORDS = {"the","and","or","of","to","a","an","is","are","was","were","be","been","in","on","for","from","with","by","as","at","it","this","that","these","those","what","why","how","does","do","can","about","into","than","their","which","who","when","where","if","please","explain","tell","show","picture","image","slide","page"}
_CHUNKS = None


def _tokenize(text):
    return {w for w in re.findall(r"[a-zA-Z0-9]+", text.lower()) if len(w) > 1 and w not in STOP_WORDS}


def _save_image(data, source_tag, image_index, ext="png"):
    image_dir = MODULE_FOLDER / ".alchemyz_image_cache"
    image_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "_", source_tag)
    path = image_dir / f"{safe}_{image_index}.{ext}"
    if not path.exists():
        path.write_bytes(data)
    return str(path)


def _extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is missing. Run: pip install pypdf") from exc
    reader = PdfReader(str(path))
    output = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        images = []
        # pypdf exposes embedded images on many, but not all, PDFs.
        try:
            for idx, image in enumerate(page.images, start=1):
                try:
                    image_data = image.data
                    image_name = getattr(image, "name", "")
                    ext = Path(image_name).suffix.lower().lstrip(".") or "png"
                    if ext not in {"png", "jpg", "jpeg", "webp"}:
                        ext = "png"
                    images.append(_save_image(image_data, f"{path.stem}_page{page_number}", idx, ext))
                except Exception:
                    continue
        except Exception:
            pass
        if text.strip() or images:
            output.append({"source": f"{path.name} - page {page_number}", "text": text, "images": images})
    return output


def _extract_pptx(path):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is missing. Run: pip install python-pptx") from exc
    presentation = Presentation(str(path))
    output = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts, images = [], []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                txt = shape.text.strip()
                if txt:
                    parts.append(txt)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text.strip())
            if getattr(shape, "shape_type", None) == 13:  # picture
                try:
                    blob = shape.image.blob
                    ext = shape.image.ext or "png"
                    images.append(_save_image(blob, f"{path.stem}_slide{slide_number}", len(images)+1, ext))
                except Exception:
                    pass
        text = "\n".join(parts).strip()
        if text or images:
            output.append({"source": f"{path.name} - slide {slide_number}", "text": text, "images": images})
    return output


def _extract_file(path):
    if not path.exists():
        return []
    if path.suffix.lower() == ".pdf":
        return _extract_pdf(path)
    if path.suffix.lower() == ".pptx":
        return _extract_pptx(path)
    return []


def _chunks_for_page(page, max_words=180):
    text = re.sub(r"\s+", " ", page["text"]).strip()
    words = text.split()
    chunks = []
    if not words:
        # Keep image-only slides/pages searchable by source title; they can be retrieved by image-related questions.
        if page["images"]:
            chunks.append({"source": page["source"], "text": "Visual content/image from chemistry learning module. " + page["source"], "images": page["images"], "tokens": _tokenize(page["source"] + " chemistry diagram image picture illustration" )})
        return chunks
    for start in range(0, len(words), max_words):
        part = " ".join(words[start:start+max_words]).strip()
        if len(part) >= 20:
            chunks.append({"source": page["source"], "text": part, "images": page["images"], "tokens": _tokenize(part + " " + page["source"])})
    return chunks


def load_module_chunks(force_reload=False):
    global _CHUNKS
    if _CHUNKS is not None and not force_reload:
        return _CHUNKS
    chunks = []
    global MODULE_FOLDER
    MODULE_FOLDER = _resolve_module_folder()
    if not MODULE_FOLDER.exists():
        print(f"[Module warning] Modules folder not found. Create a folder named 'Modules' beside Zach.py and module_search.py, then place your PDF/PPTX lesson files inside it. Expected location: {MODULE_FOLDER}")
        _CHUNKS = []
        return _CHUNKS
    files = sorted([*MODULE_FOLDER.glob("*.pdf"), *MODULE_FOLDER.glob("*.pptx")])
    if not files:
        print(f"[Module warning] No PDF or PPTX lesson files found in: {MODULE_FOLDER}")
    for path in files:
        try:
            for page in _extract_file(path):
                chunks.extend(_chunks_for_page(page))
        except Exception as error:
            print(f"[Module warning] Could not read {path.name}: {error}")
    _CHUNKS = chunks
    print(f"[Module Knowledge] Loaded {len(chunks)} text/visual sections from {len(files)} PDF/PPTX file(s).")
    return _CHUNKS


def _score_chunk(question, chunk):
    q_tokens = _tokenize(question)
    q_lower = question.lower()
    visual_intent = any(word in q_lower for word in ("picture", "image", "diagram", "illustration", "photo", "shown", "look at", "what is in", "who is this", "who is in", "explain this"))
    if not q_tokens:
        return 3 if visual_intent and chunk["images"] else 0
    overlap = q_tokens.intersection(chunk["tokens"])
    score = len(overlap)
    if visual_intent and chunk["images"]:
        score += 3
    text_lower = chunk["text"].lower()
    # More useful retrieval for specific phrases and visual references.
    for phrase in ("lewis dot", "closed system", "open system", "isolated system", "theoretical yield", "molar mass", "chemical equation", "chemical reaction", "picture", "image", "diagram", "illustration", "photo", "shown", "in this slide"):
        if phrase in q_lower and phrase in text_lower:
            score += 3
    if chunk["images"] and any(word in q_lower for word in ("picture", "image", "diagram", "illustration", "photo", "shown", "look", "visual", "what is in")):
        score += 2
    return score


def search_modules(question, top_k=4, minimum_score=1):
    scored = []
    for chunk in load_module_chunks():
        score = _score_chunk(question, chunk)
        if score >= minimum_score:
            scored.append({"source": chunk["source"], "text": chunk["text"], "images": chunk["images"], "score": score})
    scored.sort(key=lambda item: item["score"], reverse=True)
    # Avoid returning repeated chunks from the same slide/page where possible.
    selected, seen = [], set()
    for item in scored:
        if item["source"] not in seen:
            selected.append(item)
            seen.add(item["source"])
        if len(selected) >= top_k:
            break
    return selected


def build_module_context(question, top_k=4, minimum_score=1):
    results = search_modules(question, top_k=top_k, minimum_score=minimum_score)
    if not results:
        return "", []
    context = "\n\n".join(f"[Module source: {r['source']}]\n{r['text']}" for r in results)
    return context, results


def module_supports_question(question, minimum_score=1):
    results = search_modules(question, top_k=4, minimum_score=minimum_score)
    return bool(results), results


if __name__ == "__main__":
    print("AlchemyZ module search test")
    print(f"Folder: {MODULE_FOLDER}")
    load_module_chunks()
    while True:
        question = input("Question (or exit): ").strip()
        if question.lower() == "exit":
            break
        for item in search_modules(question):
            print(f"\n{item['source']} | score={item['score']} | images={len(item['images'])}")
            print(item['text'][:450])
