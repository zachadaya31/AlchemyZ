import re
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation


# ============================================================
# ALCHEMYZ MODULE FOLDER
# ============================================================

MODULE_FOLDER = Path(__file__).parent / "Modules"


# ============================================================
# SEARCH SETTINGS
# ============================================================

DEFAULT_TOP_K = 4
DEFAULT_MINIMUM_SCORE = 2


# ============================================================
# STOP WORDS
# ============================================================

STOP_WORDS = {
    "a",
    "an",
    "and",
    "are",
    "be",
    "can",
    "do",
    "does",
    "for",
    "how",
    "i",
    "in",
    "is",
    "it",
    "me",
    "of",
    "on",
    "or",
    "the",
    "this",
    "that",
    "to",
    "was",
    "what",
    "when",
    "where",
    "which",
    "who",
    "why",
    "with",
    "you",
    "your",
    "tell",
    "explain",
    "please",
}


# ============================================================
# GET MODULE FILES
# ============================================================

def get_module_files():

    if not MODULE_FOLDER.exists():

        print(
            f"WARNING: Modules folder was not found:\n"
            f"{MODULE_FOLDER}"
        )

        return []

    files = []

    files.extend(
        MODULE_FOLDER.glob("*.pdf")
    )

    files.extend(
        MODULE_FOLDER.glob("*.pptx")
    )

    return sorted(files)


# ============================================================
# EXTRACT PDF TEXT
# ============================================================

def extract_pdf_text(file_path):

    pages = []

    try:

        reader = PdfReader(
            str(file_path)
        )

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            page_text = page.extract_text()

            if page_text:

                pages.append(
                    f"[{file_path.name} - Page {page_number}]\n"
                    f"{page_text}"
                )

    except Exception as e:

        print(
            f"Could not read PDF "
            f"{file_path.name}: {e}"
        )

    return "\n\n".join(pages)


# ============================================================
# EXTRACT POWERPOINT TEXT
# ============================================================

def extract_pptx_text(file_path):

    slides_text = []

    try:

        presentation = Presentation(
            str(file_path)
        )

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_parts = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:

                        slide_parts.append(
                            text
                        )

            if slide_parts:

                slides_text.append(
                    f"[{file_path.name} - "
                    f"Slide {slide_number}]\n"
                    +
                    "\n".join(slide_parts)
                )

    except Exception as e:

        print(
            f"Could not read PPTX "
            f"{file_path.name}: {e}"
        )

    return "\n\n".join(slides_text)


# ============================================================
# EXTRACT MODULE TEXT
# ============================================================

def extract_module(file_path):

    extension = file_path.suffix.lower()

    if extension == ".pdf":

        return extract_pdf_text(
            file_path
        )

    elif extension == ".pptx":

        return extract_pptx_text(
            file_path
        )

    return ""


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# TOKENIZE
# ============================================================

def tokenize(text):

    words = re.findall(
        r"[a-zA-Z0-9]+",
        text.lower()
    )

    return [
        word
        for word in words
        if word not in STOP_WORDS
    ]


# ============================================================
# NORMALIZE PHRASE
# ============================================================

def normalize_phrase(text):

    text = text.lower()

    text = re.sub(
        r"[^a-z0-9\s]",
        " ",
        text
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# CHUNK TEXT
# ============================================================

def chunk_text(
    text,
    chunk_size=180
):

    words = text.split()

    chunks = []

    for i in range(
        0,
        len(words),
        chunk_size
    ):

        chunk = " ".join(
            words[
                i:i + chunk_size
            ]
        )

        if chunk.strip():

            chunks.append(
                chunk.strip()
            )

    return chunks


# ============================================================
# LOAD ALL MODULES
# ============================================================

def load_modules():

    module_files = get_module_files()

    all_chunks = []

    for file_path in module_files:

        text = extract_module(
            file_path
        )

        if not text:

            continue

        text = clean_text(
            text
        )

        chunks = chunk_text(
            text
        )

        for chunk in chunks:

            all_chunks.append(
                {
                    "file": file_path.name,
                    "text": chunk
                }
            )

    return all_chunks


# ============================================================
# SCORE MODULE CHUNK
# ============================================================

def score_chunk(
    question,
    chunk
):

    question_normalized = normalize_phrase(
        question
    )

    chunk_normalized = normalize_phrase(
        chunk
    )

    question_words = tokenize(
        question
    )

    chunk_words = set(
        tokenize(chunk)
    )

    if not question_words:

        return 0


    score = 0


    # --------------------------------------------------------
    # EXACT PHRASE MATCH
    # --------------------------------------------------------

    if question_normalized in chunk_normalized:

        score += 10


    # --------------------------------------------------------
    # WORD OVERLAP
    # --------------------------------------------------------

    overlap = set(
        question_words
    ).intersection(
        chunk_words
    )

    score += len(overlap)


    # --------------------------------------------------------
    # EXTRA WEIGHT FOR MULTIPLE MATCHES
    # --------------------------------------------------------

    if len(overlap) >= 2:

        score += 2

    if len(overlap) >= 3:

        score += 2


    return score


# ============================================================
# SEARCH MODULES
#
# IMPORTANT:
# minimum_score is handled INSIDE this function.
# ============================================================

def search_modules(
    question,
    top_k=DEFAULT_TOP_K,
    minimum_score=DEFAULT_MINIMUM_SCORE
):

    modules = load_modules()

    if not modules:

        return []


    results = []


    for module in modules:

        score = score_chunk(
            question,
            module["text"]
        )

        if score >= minimum_score:

            results.append(
                {
                    "file": module["file"],
                    "text": module["text"],
                    "score": score
                }
            )


    results.sort(
        key=lambda item: item["score"],
        reverse=True
    )


    return results[:top_k]


# ============================================================
# BUILD MODULE CONTEXT
# ============================================================

def build_module_context(
    question,
    top_k=DEFAULT_TOP_K
):

    results = search_modules(
        question,
        top_k=top_k,
        minimum_score=DEFAULT_MINIMUM_SCORE
    )


    if not results:

        return "", []


    context_parts = []


    for result in results:

        context_parts.append(
            f"SOURCE: {result['file']}\n"
            f"{result['text']}"
        )


    context = "\n\n".join(
        context_parts
    )


    return context, results


# ============================================================
# CHECK WHETHER MODULES SUPPORT QUESTION
# ============================================================

def module_supports_question(
    question
):

    results = search_modules(
        question,
        top_k=DEFAULT_TOP_K,
        minimum_score=DEFAULT_MINIMUM_SCORE
    )


    if not results:

        return False, []


    return True, results


# ============================================================
# TEST MODULE SEARCH
# ============================================================

if __name__ == "__main__":

    print()
    print(
        "======================================"
    )
    print(
        "       ALCHEMYZ MODULE SEARCH"
    )
    print(
        "======================================"
    )
    print()

    print(
        f"Module folder:"
    )

    print(
        MODULE_FOLDER
    )

    print()


    files = get_module_files()


    print(
        f"Found {len(files)} module file(s)."
    )

    print()


    if not files:

        print(
            "No PDF or PPTX files were found."
        )

    else:

        for file in files:

            print(
                f"- {file.name}"
            )


    print()


    question = input(
        "Enter a chemistry question: "
    )


    context, results = build_module_context(
        question
    )


    print()
    print(
        "Relevant module sections:"
    )
    print()


    if not results:

        print(
            "No relevant module sections found."
        )

    else:

        for result in results:

            print(
                f"SOURCE: {result['file']}"
            )

            print(
                f"SCORE: {result['score']}"
            )

            print(
                result["text"]
            )

            print()

            print(
                "--------------------------------------"
            )

            print()